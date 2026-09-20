"""可编辑性回读验证与项目导出接口。"""

from __future__ import annotations

import uuid
from collections.abc import AsyncGenerator
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.api.deps import get_queue
from app.core.config import get_settings
from app.core.db import async_session_factory
from app.domain.content import BulletsBlock, Deck, Slide, TextBlock
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT
from app.domain.sample import load_sample_deck
from app.main import app
from app.models.project import Project, ProjectOutline, ProjectSource
from app.models.slide import Slide as SlideRow
from app.render.pptx import render_deck_to_pptx
from app.render.verify import verify_pptx
from app.services.preview import preview_key
from app.storage import get_storage
from app.worker.preview_tasks import generate_preview

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
async def client() -> AsyncGenerator[AsyncClient, None]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as value:
        yield value


async def _sign_up(client: AsyncClient) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/register",
        json={"email": f"export_{uuid.uuid4().hex}@example.com", "password": "password123"},
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def _seed_ready_project(
    *,
    title: str = "导出测试",
    blocks: list[dict] | None = None,
    status: str = "ready",
    theme_id: str = "ivory",
) -> tuple[str, dict[str, str]]:
    """直接落库一份可导出的单页项目，绕过生成队列。"""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        headers = await _sign_up(client)
        created = await client.post(
            "/api/v1/projects",
            json={"title": title, "page_count": 5, "theme_id": theme_id},
            headers=headers,
        )
        assert created.status_code == 201, created.text
        project_id = created.json()["id"]

    page_id = uuid.uuid4()
    slide_blocks = blocks or [
        {"id": "t1", "slot_id": "title", "type": "text", "text": "增长复盘", "locked": False},
        {
            "id": "b1",
            "slot_id": "body",
            "type": "bullets",
            "items": ["营收增长 37%", "用户破千万"],
            "locked": False,
        },
    ]

    async with async_session_factory() as session:
        result = await session.execute(
            select(Project)
            .options(selectinload(Project.outline), selectinload(Project.sources))
            .where(Project.id == uuid.UUID(project_id))
        )
        record = result.scalar_one()
        session.add(
            ProjectSource(
                project_id=record.id,
                kind="text",
                sections=[
                    {
                        "level": 0,
                        "heading": None,
                        "text": "营收增长 37%，用户破千万",
                        "locator": "第 1 段",
                    }
                ],
                warnings=[],
                char_count=20,
            )
        )
        session.add(
            ProjectOutline(
                project_id=record.id,
                status="confirmed",
                pages=[
                    {
                        "id": str(page_id),
                        "title": "增长复盘",
                        "objective": "目标",
                        "key_points": ["要点一", "要点二"],
                        "source_refs": ["S1:1"],
                        "layout_id": "bullets",
                    }
                ],
                revision=2,
            )
        )
        session.add(
            SlideRow(
                project_id=record.id,
                outline_page_id=page_id,
                position=1,
                layout_id="bullets",
                layout_mode="fixed",
                title="增长复盘",
                status=status,
                blocks=slide_blocks if status == "ready" else [],
                issues=[],
                revision=1,
            )
        )
        record.status = "ready" if status == "ready" else "generating"
        await session.commit()

    return project_id, headers


# --- 验证模块 ---


def test_verify_sample_deck_passes() -> None:
    deck = load_sample_deck()
    payload = render_deck_to_pptx(deck).getvalue()
    report = verify_pptx(payload, deck)

    assert report.passed is True, [issue.message for issue in report.issues]
    assert report.slide_count == len(deck.slides)
    assert report.expected_slide_count == len(deck.slides)
    assert report.issues == []


def test_verify_detects_page_count_mismatch() -> None:
    deck = load_sample_deck()
    short = deck.model_copy(update={"slides": deck.slides[:1]})
    payload = render_deck_to_pptx(short).getvalue()

    report = verify_pptx(payload, deck)
    assert report.passed is False
    assert any(issue.check == "page_count" for issue in report.issues)


def test_verify_detects_out_of_bounds_shape() -> None:
    deck = Deck(
        id="d1",
        title="越界",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    slide = presentation.slides[0]
    # 故意放到画布外
    slide.shapes.add_textbox(
        Emu(Pt(CANVAS_WIDTH_PT + 40)),
        Emu(Pt(10)),
        Emu(Pt(100)),
        Emu(Pt(40)),
    )
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "bounds" and issue.slide_index == 1 for issue in report.issues)


def test_verify_detects_missing_text() -> None:
    deck = Deck(
        id="d1",
        title="缺字",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="真实标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["真实要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    # 清空文本框，制造「文字丢失」
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            shape.text_frame.clear()
    bad = BytesIO()
    presentation.save(bad)

    # 源 Deck 仍要求这些文字存在
    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check in {"text_frame", "content_integrity"} for issue in report.issues)


def test_verify_detects_full_page_picture() -> None:
    deck = Deck(
        id="d1",
        title="整页图",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    buffer = render_deck_to_pptx(deck)
    presentation = Presentation(buffer)
    # 1x1 PNG
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00"
        b"\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    presentation.slides[0].shapes.add_picture(
        BytesIO(png),
        Emu(0),
        Emu(0),
        Emu(Pt(CANVAS_WIDTH_PT)),
        Emu(Pt(CANVAS_HEIGHT_PT)),
    )
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "full_page_picture" for issue in report.issues)


def test_verify_detects_wrong_canvas_size() -> None:
    deck = Deck(
        id="d1",
        title="尺寸",
        theme_id="ivory",
        slides=[
            Slide(
                id="s1",
                layout_id="bullets",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    BulletsBlock(id="b1", slot_id="body", items=["要点"]),
                ],
            )
        ],
    )
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    bad = BytesIO()
    presentation.save(bad)

    report = verify_pptx(bad.getvalue(), deck)
    assert report.passed is False
    assert any(issue.check == "canvas_size" for issue in report.issues)


# --- 导出接口 ---


@pytest.mark.asyncio
async def test_export_endpoint_returns_pptx(client: AsyncClient) -> None:
    project_id, headers = await _seed_ready_project(title="中文文件名导出")
    response = await client.get(f"/api/v1/projects/{project_id}/deck/export", headers=headers)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(PPTX_MEDIA_TYPE.split(";")[0])
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.content[:2] == b"PK"

    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides) == 1
    assert presentation.slide_width == Pt(CANVAS_WIDTH_PT)
    assert presentation.slide_height == Pt(CANVAS_HEIGHT_PT)


@pytest.mark.asyncio
async def test_export_blocked_when_pages_incomplete(client: AsyncClient) -> None:
    project_id, headers = await _seed_ready_project(status="pending")
    response = await client.get(f"/api/v1/projects/{project_id}/deck/export", headers=headers)

    assert response.status_code == 409
    assert response.json()["detail"] == "页面尚未全部生成完成，无法导出"


@pytest.mark.asyncio
async def test_export_blocked_on_quality_errors(client: AsyncClient) -> None:
    # 缺必填 body 槽位 → error
    project_id, headers = await _seed_ready_project(
        blocks=[
            {"id": "t1", "slot_id": "title", "type": "text", "text": "仅标题", "locked": False},
        ]
    )
    response = await client.get(f"/api/v1/projects/{project_id}/deck/export", headers=headers)

    assert response.status_code == 409, response.text
    detail = response.json()["detail"]
    assert detail["message"] == "导出前检查未通过，存在必须修复的问题"
    assert detail["report"]["export_allowed"] is False
    assert any(issue["severity"] == "error" for issue in detail["report"]["issues"])


@pytest.mark.asyncio
async def test_export_allows_warnings(client: AsyncClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from pathlib import Path

    from app.domain import text_metrics as tm

    # 字体缺失只产生 warning，不应阻断
    monkeypatch.setattr(tm, "FONTS_DIR", Path("/tmp/aippt-no-fonts-export"))
    tm.clear_font_cache()

    project_id, headers = await _seed_ready_project()
    response = await client.get(f"/api/v1/projects/{project_id}/deck/export", headers=headers)

    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"


async def test_preview_worker_export_and_owner_isolation(client, monkeypatch):
    project_id, headers = await _seed_ready_project()
    settings = get_settings()
    monkeypatch.setattr(settings, "preview_soffice", "test-soffice")
    monkeypatch.setattr(settings, "export_require_preview", True)
    monkeypatch.setattr(
        "app.worker.preview_tasks.render_pptx_preview", lambda *a, **k: [b"png-data"]
    )
    url = f"/api/v1/projects/{project_id}/deck"
    before = await client.get(f"{url}/export", headers=headers)
    assert before.status_code == 409
    initial = (await client.get(f"{url}/preview", headers=headers)).json()
    fingerprint = initial["fingerprint"]
    assert initial["status"] == "missing"
    await generate_preview({}, project_id, fingerprint)
    ready = (await client.get(f"{url}/preview", headers=headers)).json()
    assert ready["status"] == "ready"
    png = await client.get(f"{url}/preview/{fingerprint}/pages/1", headers=headers)
    assert png.content == b"png-data"
    other = await _sign_up(client)
    denied = await client.get(f"{url}/preview/{fingerprint}/pages/1", headers=other)
    assert denied.status_code == 404
    exported = await client.get(f"{url}/export", headers=headers)
    assert exported.status_code == 200
    async with async_session_factory() as session:
        project = await session.get(Project, uuid.UUID(project_id))
        cached = get_storage().load(
            f"{preview_key(project.user_id, project.id, fingerprint)}/deck.pptx"
        )
        assert exported.content == cached
        project.theme_id = "enterprise-dark"
        await session.commit()
    changed = (await client.get(f"{url}/preview", headers=headers)).json()
    assert changed["fingerprint"] != fingerprint
    assert changed["status"] == "missing"
    assert (await client.get(f"{url}/export", headers=headers)).status_code == 409


async def test_preview_queue_failure_is_retryable(client, monkeypatch):
    project_id, headers = await _seed_ready_project()
    monkeypatch.setattr(get_settings(), "preview_soffice", "test-soffice")

    class Queue:
        async def enqueue_job(self, *args, **kwargs):
            return None

    app.dependency_overrides[get_queue] = Queue
    try:
        url = f"/api/v1/projects/{project_id}/deck/preview"
        response = await client.post(url, headers=headers)
        assert response.status_code == 503
        assert (await client.get(url, headers=headers)).json()["status"] == "failed"
    finally:
        app.dependency_overrides.pop(get_queue, None)


async def test_preview_render_failure_does_not_claim_ready(client, monkeypatch):
    project_id, headers = await _seed_ready_project()
    monkeypatch.setattr(get_settings(), "preview_soffice", "test-soffice")

    def fail(*args, **kwargs):
        raise RuntimeError("engine failed")

    monkeypatch.setattr("app.worker.preview_tasks.render_pptx_preview", fail)
    url = f"/api/v1/projects/{project_id}/deck/preview"
    fingerprint = (await client.get(url, headers=headers)).json()["fingerprint"]
    await generate_preview({}, project_id, fingerprint)
    result = (await client.get(url, headers=headers)).json()
    assert result["status"] == "failed"
    assert result["page_count"] == 0
