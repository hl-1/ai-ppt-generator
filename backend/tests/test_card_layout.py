from copy import deepcopy

import pytest
from pptx import Presentation

from app.domain.block_style import BlockStyle
from app.domain.card_layout import card_grid, cards_natural_height
from app.domain.content import CardItem, CardsBlock, Deck, Slide, TextBlock
from app.domain.flex_layout import FlexContainer, FlexLeaf
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT, SAFE_AREA_WIDTH_PT
from app.domain.slide_geometry import placed_by_block_id
from app.domain.theme import get_theme, load_themes, resolve_theme
from app.domain.theme_layout import reflow_slide_to_theme, theme_requires_reflow
from app.domain.validation import validate_slide
from app.models.slide import Slide as SlideRow
from app.render.pptx import render_deck_to_pptx
from app.services.deck import refresh_slide_issues

THEMES = list(load_themes())


def cards_slide(count=6):
    return Slide(
        id="cards-page",
        layout_id="bullets",
        layout_mode="flex",
        layout_tree=FlexContainer(
            type="column",
            id="root",
            gap_pt=16,
            children=[
                FlexLeaf(id="title-leaf", block_id="title", text_style="title", grow=1),
                FlexLeaf(id="cards-leaf", block_id="cards", grow=0.5),
            ],
        ),
        blocks=[
            TextBlock(id="title", slot_id="title", text="工作中的技术提升"),
            CardsBlock(
                id="cards",
                slot_id="body",
                locked=True,
                items=[
                    CardItem(
                        title=f"练习任务 {index + 1}",
                        desc="在真实任务中练习，结合测试结果记录反馈并改进。",
                    )
                    for index in range(count)
                ],
            ),
        ],
    )


@pytest.mark.parametrize("theme_id", THEMES)
def test_multirow_cards_fit_and_export_with_the_same_grid(theme_id):
    theme = get_theme(theme_id)
    original = cards_slide()
    fitted = reflow_slide_to_theme(original, theme)
    assert fitted.blocks == original.blocks
    assert not [issue for issue in validate_slide(fitted, theme=theme) if issue.code == "overflow"]
    rect = placed_by_block_id(fitted)["cards"].rect
    columns, width, height = card_grid(
        6, rect.w * CANVAS_WIDTH_PT, rect.h * CANVAS_HEIGHT_PT, visual_style=theme.visual_style
    )
    deck = Deck(id="theme-qa", title="主题检查", theme_id=theme_id, slides=[fitted])
    pptx = Presentation(render_deck_to_pptx(deck))
    panels = [
        shape
        for shape in pptx.slides[0].shapes
        if shape.name.startswith(("card-panel-", "financial-panel-"))
    ]
    assert len(panels) == 6
    assert float(panels[0].width.pt) == pytest.approx(width, abs=0.01)
    assert float(panels[0].height.pt) == pytest.approx(height, abs=0.01)
    assert panels[columns].top > panels[0].top
    exported = "\n".join(shape.text for shape in pptx.slides[0].shapes if shape.has_text_frame)
    for item in original.blocks[1].items:
        assert item.title in exported
        assert item.desc in exported


@pytest.mark.parametrize("theme_id", THEMES)
def test_card_height_accounts_for_width_font_and_rows(theme_id):
    theme = get_theme(theme_id)
    block = cards_slide().blocks[1]
    single_row = cards_natural_height(
        block.model_copy(update={"items": block.items[:3]}), theme, SAFE_AREA_WIDTH_PT
    )
    multi_row = cards_natural_height(block, theme, SAFE_AREA_WIDTH_PT)
    assert multi_row >= 2 * single_row
    narrow = cards_natural_height(block, theme, 260)
    assert narrow > multi_row
    large = block.model_copy(update={"style": BlockStyle(size_pt=24)})
    assert cards_natural_height(large, theme, 520) > cards_natural_height(block, theme, 520)


def test_theme_reflow_updates_layout_revision_without_changing_locked_content():
    source = cards_slide()
    row = SlideRow(
        layout_mode="flex",
        layout_id=source.layout_id,
        revision=3,
        layout_tree=source.layout_tree.model_dump(mode="json"),
        blocks=[block.model_dump(mode="json") for block in source.blocks],
    )
    before = deepcopy(row.blocks)
    refresh_slide_issues(row, theme=get_theme("enterprise"), reflow=True)
    assert row.revision == 4
    assert row.blocks == before
    assert row.layout_tree != source.layout_tree.model_dump(mode="json")
    assert not [issue for issue in row.issues if issue.get("code") == "overflow"]


def test_palette_changes_do_not_trigger_layout_reflow():
    original = get_theme("enterprise")
    recolored = resolve_theme("enterprise", {"palette": {"accent": "#123456"}})
    assert not theme_requires_reflow(original, recolored)
    assert theme_requires_reflow(original, get_theme("financial-tech"))
    larger = resolve_theme("enterprise", {"text_styles": {"body": {"size_pt": 24}}})
    assert theme_requires_reflow(original, larger)


def test_empty_cards_and_padding_are_measured_safely():
    empty = cards_slide(1).blocks[1].model_copy(update={"items": []})
    theme = get_theme("ivory")
    assert cards_natural_height(empty, theme, 0) == 0
    assert card_grid(0, 0, 0) == (1, 1, 1)
    block = cards_slide(3).blocks[1]
    padded = block.model_copy(
        update={"style": BlockStyle(padding_pt=18, border_width_pt=2, border_color="line")}
    )
    assert cards_natural_height(padded, theme, 540) > cards_natural_height(block, theme, 540)


@pytest.mark.parametrize("theme_id", THEMES)
def test_impossible_card_content_is_reported_instead_of_passing_validation(theme_id):
    theme = get_theme(theme_id)
    source = cards_slide()
    tree = source.layout_tree.model_copy(
        update={
            "children": [
                source.layout_tree.children[0].model_copy(update={"grow": 10}),
                source.layout_tree.children[1].model_copy(update={"grow": 0.1}),
            ]
        }
    )
    source = source.model_copy(update={"layout_tree": tree})
    assert any(issue.code == "overflow" for issue in validate_slide(source, theme=theme))
