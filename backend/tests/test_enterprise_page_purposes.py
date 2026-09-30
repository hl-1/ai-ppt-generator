from collections import Counter

import pytest
from pptx import Presentation

from app.domain.content import (
    BulletsBlock,
    CardsBlock,
    ChartBlock,
    ComboChartBlock,
    Deck,
    FinancialTableBlock,
    KpiBlock,
    Slide,
    TextBlock,
)
from app.domain.enterprise_layout import bind_planned_evidence, compose_report_slide
from app.domain.enterprise_quality import check_planned_data
from app.domain.evidence import prepare_page_plan
from app.domain.flex_fit import fit_tree_to_content
from app.domain.flex_layout import FlexContainer, FlexLeaf, iter_leaf_block_ids
from app.domain.flex_width import fit_row_widths
from app.domain.outline import EvidenceItem, OutlinePageDraft
from app.domain.page_rhythm import assign_layout_templates
from app.domain.quality import check_evidence_alignment
from app.domain.slide_geometry import placed_by_block_id
from app.domain.theme import get_theme
from app.domain.topic_visuals import normalize_topic_pages
from app.domain.validation import validate_slide
from app.render.pptx import render_deck_to_pptx


def trend_evidence():
    return [
        EvidenceItem(
            source_ref="S1:1",
            metric=metric,
            unit=unit,
            scope="同一业务范围",
            period=period,
            value=str(value),
            quote=f"{period}同一业务范围{metric}{value}{unit}。",
        )
        for metric, unit, values in [
            ("成熟度", "分", [52, 61, 70, 78]),
            ("采用率", "%", [18, 31, 47, 63]),
        ]
        for period, value in zip(["Q1", "Q2", "Q3", "Q4"], values, strict=True)
    ]


def outline(**changes):
    values = dict(
        title="成熟度与采用率双线提升",
        objective="比较两项指标的全年变化",
        key_points=["成熟度从52分提升至78分", "采用率从18%提升至63%", "下一阶段推进试点验证"],
        layout_id="bullets",
        narrative_role="performance",
        evidence_kind="trend",
        visual_type="line",
        key_message="两项指标持续提升，后续应关注落地差异",
        evidence=trend_evidence(),
    )
    return OutlinePageDraft(**(values | changes))


def draft(page, extra=None):
    blocks = [TextBlock(id="title", slot_id="title", text=page.title), *(extra or [])]
    return Slide(
        id="verification",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )


def fitted(slide, page):
    theme = get_theme("enterprise")
    wide = fit_row_widths(slide.layout_tree, slide.blocks, theme=theme)
    return slide.model_copy(
        update={
            "layout_tree": fit_tree_to_content(
                wide,
                slide.blocks,
                theme=theme,
                page_role=page.page_role,
            )
        }
    )


def test_dual_trends_preserve_both_metrics_and_independent_units():
    page = outline()
    sources = {"S1:1": "\n".join(item.quote for item in page.evidence)}
    prepared = prepare_page_plan(page, sources)
    assert prepared.evidence_kind == "trend"
    assert len(prepared.evidence) == 8
    result = fitted(compose_report_slide(bind_planned_evidence(draft(page), page), page), page)
    charts = [block for block in result.blocks if isinstance(block, ChartBlock)]

    assert len(charts) == 2
    assert [(chart.series[0].name, chart.unit) for chart in charts] == [
        ("成熟度", "分"),
        ("采用率", "%"),
    ]
    assert charts[0].series[0].values == [52, 61, 70, 78]
    assert charts[1].series[0].values == [18, 31, 47, 63]
    assert check_planned_data(result, page) == []
    assert not any(
        issue.severity == "error" for issue in validate_slide(result, theme=get_theme("enterprise"))
    )
    placements = placed_by_block_id(result)
    assert placements[charts[0].id].rect.x < placements[charts[1].id].rect.x
    missing = result.model_copy(
        update={"blocks": [block for block in result.blocks if block is not charts[1]]}
    )
    assert any(
        issue.code == "evidence_metric_missing" for issue in check_planned_data(missing, page)
    )


@pytest.mark.parametrize("role", ["executive_summary", "summary", "decision", "risk", "action"])
def test_business_roles_do_not_receive_unrelated_auto_charts(role):
    page = outline(narrative_role=role, visual_type="auto", evidence_kind="narrative", evidence=[])
    normalized = normalize_topic_pages([page], {})[0]
    assert normalized.visual_type == "auto"
    assert normalized.evidence == []


def test_same_selected_visual_is_kept_on_pages_with_different_subjects():
    first = outline()
    second = outline(title="另一业务范围的变化", objective="分析另一项指标")
    normalized = normalize_topic_pages([first, second], {})
    assert [page.visual_type for page in normalized] == ["line", "line"]
    assert all(len(page.evidence) == 8 for page in normalized)


def test_summary_shows_latest_results_and_decision_request():
    page = outline(
        title="总结与决策请求",
        page_role="summary",
        narrative_role="summary",
        visual_type="auto",
        evidence_kind="narrative",
    )
    model_chart = ChartBlock(
        id="irrelevant-chart",
        slot_id="visual",
        chart_type="bar",
        categories=["A", "B"],
        series=[{"name": "无关资源", "values": [38, 27]}],
    )
    bound = bind_planned_evidence(draft(page, [model_chart]), page)
    result = fitted(
        compose_report_slide(bound, page, decision_request="确认下一阶段试点方向"), page
    )
    metrics = [block for block in result.blocks if isinstance(block, KpiBlock)]

    assert [(metric.label, metric.value) for metric in metrics] == [
        ("成熟度", "78分"),
        ("采用率", "63%"),
    ]
    assert any(isinstance(block, BulletsBlock) for block in result.blocks)
    assert any(
        block.type == "text" and "待决策：确认下一阶段试点方向" in block.text
        for block in result.blocks
    )
    assert not any(block.type == "chart" for block in result.blocks)
    assert not check_planned_data(result, page)


def test_agenda_has_title_on_top_and_wide_numbered_rows():
    page = outline(
        title="汇报目录",
        page_role="toc",
        visual_type="auto",
        evidence=[],
        evidence_kind="narrative",
    )
    cards = CardsBlock(
        id="long-agenda",
        slot_id="cards",
        items=[
            {"title": title, "desc": "说明相关业务判断、指标变化、影响范围及后续推进事项。" * 3}
            for title in ["执行摘要", "表现与趋势", "资源结构", "推进路径", "风险与行动"]
        ],
    )
    result = fitted(
        compose_report_slide(draft(page, [cards]), page, layout_template="opening_split_bottom"),
        page,
    )
    placements = placed_by_block_id(result)
    title_bottom = placements["title"].rect.y + placements["title"].rect.h
    headings = [block for block in result.blocks if block.slot_id == "agenda_heading"]

    assert len(headings) == 5
    assert all(placements[block.id].rect.y >= title_bottom for block in headings)
    assert all(placements[block.id].rect.w > 0.3 for block in headings)
    assert not any(block.type == "cards" for block in result.blocks)
    assert Counter(iter_leaf_block_ids(result.layout_tree)) == Counter(
        block.id for block in result.blocks
    )
    assert "未放入画布" in result.speaker_notes
    assert assign_layout_templates([page], seed="agenda", topic_mode=True)[1].id == "agenda_list"


def test_same_unit_trends_share_one_chart_and_preserve_all_series():
    page = outline(evidence=[item.model_copy(update={"unit": "分"}) for item in trend_evidence()])
    result = fitted(compose_report_slide(bind_planned_evidence(draft(page), page), page), page)
    charts = [block for block in result.blocks if isinstance(block, ChartBlock)]
    assert len(charts) == 1
    assert [series.name for series in charts[0].series] == ["成熟度", "采用率"]
    assert check_planned_data(result, page) == []


def test_summary_can_express_trend_results_as_verified_kpis():
    page = outline(page_role="summary", narrative_role="summary", visual_type="auto")
    result = fitted(compose_report_slide(bind_planned_evidence(draft(page), page), page), page)
    assert check_planned_data(result, page) == []
    assert check_evidence_alignment(result, "trend", summary_metrics=True) == []
    missing = result.model_copy(
        update={
            "blocks": [
                block
                for block in result.blocks
                if not (block.type == "kpi" and block.label == "采用率")
            ]
        }
    )
    assert any(
        issue.code == "evidence_metric_missing" for issue in check_planned_data(missing, page)
    )


def test_auto_combo_uses_both_metrics_and_passes_evidence_alignment():
    page = outline(visual_type="auto")
    sources = {
        "S1:1": "\n".join(
            f"指标：{item.metric}；分类：；数值：{item.value}；单位：{item.unit}；时间：{item.period}；范围：{item.scope}。"
            for item in page.evidence
        )
    }
    planned = normalize_topic_pages([page], sources)[0]
    assert planned.visual_type == "combo_chart"
    result = fitted(
        compose_report_slide(bind_planned_evidence(draft(planned), planned), planned), planned
    )
    combo = next(block for block in result.blocks if isinstance(block, ComboChartBlock))
    assert {series.name for series in [*combo.bars, *combo.lines]} == {"成熟度", "采用率"}
    assert check_evidence_alignment(result, "comparison") == []


def test_financial_table_is_recognized_as_table_evidence():
    slide = draft(
        outline(),
        [
            FinancialTableBlock(
                id="finance",
                slot_id="visual",
                columns=["本期", "上期"],
                rows=[{"label": "收入", "values": ["120", "100"]}],
                unit="万元",
            )
        ],
    )
    assert check_evidence_alignment(slide, "table") == []


def test_uncurated_agenda_keeps_original_blocks_without_pairing_assumption():
    page = outline(page_role="toc", evidence=[])
    source = draft(page, [BulletsBlock(id="agenda", slot_id="body", items=["表现", "行动"])])
    result = compose_report_slide(source, page, curate=False)
    assert Counter(iter_leaf_block_ids(result.layout_tree)) == Counter(
        block.id for block in result.blocks
    )


def test_export_preserves_both_native_charts_and_business_insights():
    page = outline()
    result = fitted(compose_report_slide(bind_planned_evidence(draft(page), page), page), page)
    deck = Deck(id="verification", title=page.title, theme_id="enterprise", slides=[result])
    presentation = Presentation(render_deck_to_pptx(deck))
    exported = presentation.slides[0]
    charts = [shape.chart for shape in exported.shapes if shape.has_chart]
    assert len(charts) == 2
    assert {series.name for chart in charts for series in chart.series} == {"成熟度", "采用率"}
    text = "\n".join(shape.text for shape in exported.shapes if shape.has_text_frame)
    assert "采用率从18%提升至63%" in text
