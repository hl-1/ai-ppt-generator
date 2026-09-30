from copy import deepcopy

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from app.domain.content import DiagramEdge, DiagramNode
from app.domain.financial_tech import panel_grid, platform_node_rects
from app.domain.geometry import CANVAS_HEIGHT_PT, CANVAS_WIDTH_PT, Rect
from app.domain.sample import load_financial_tech_sample, load_sample_deck
from app.domain.theme import get_theme, resolve_theme
from app.domain.validation import validate_deck
from app.main import app
from app.render.pptx import render_deck_to_pptx
from app.render.verify import verify_pptx


def test_financial_theme_overrides_keep_component_style() -> None:
    theme = resolve_theme("financial-tech", {"palette": {"accent": "#3377BB"}})
    assert theme.visual_style == "financial-tech"
    assert theme.palette.accent == "#3377BB"
    assert theme.palette.chart_series[0] == "#3377BB"
    assert get_theme("enterprise-dark").visual_style == "standard"


@pytest.mark.parametrize("count,width", [(1, 120), (4, 856), (6, 520), (9, 350)])
def test_panels_wrap_to_fit(count: int, width: float) -> None:
    columns, card_width, _ = panel_grid(count, width, 320)
    assert 1 <= columns <= 4
    assert columns * card_width + (columns - 1) * 16 <= width + 0.01


@pytest.mark.parametrize("size", [(856, 320), (300, 240), (650, 180)])
def test_platform_geometry_stays_inside_slot(size: tuple[float, float]) -> None:
    sample = load_financial_tech_sample().slides[1].blocks[1]
    rect = Rect(x=0.05, y=0.2, w=size[0] / CANVAS_WIDTH_PT, h=size[1] / CANVAS_HEIGHT_PT)
    boxes, hub = platform_node_rects(sample.nodes, sample.edges, rect)
    assert len(boxes) == len(sample.nodes)
    assert hub == ("hub" if size == (856, 320) else None)
    for box in boxes.values():
        assert rect.x <= box.x <= rect.right
        assert rect.y <= box.y <= rect.bottom
        assert box.right <= rect.right + 1e-6
        assert box.bottom <= rect.bottom + 1e-6
    values = list(boxes.values())
    for index, first in enumerate(values):
        for second in values[index + 1 :]:
            overlap_w = min(first.right, second.right) - max(first.x, second.x)
            overlap_h = min(first.bottom, second.bottom) - max(first.y, second.y)
            assert overlap_w <= 0 or overlap_h <= 0


def test_cyclic_platform_uses_grid() -> None:
    nodes = [DiagramNode(id=str(i), title=str(i)) for i in range(5)]
    edges = [DiagramEdge(source=str(i), target=str((i + 1) % 5)) for i in range(5)]
    boxes, hub = platform_node_rects(nodes, edges, Rect(x=0, y=0, w=0.8, h=0.6))
    assert hub is None
    assert len(boxes) == 5
    assert boxes["3"].y > boxes["0"].y


def test_dense_platform_fits_short_slot() -> None:
    nodes = [DiagramNode(id=str(i), title=str(i)) for i in range(18)]
    edges = [DiagramEdge(source=str(i), target=str(i + 1)) for i in range(17)]
    rect = Rect(x=0, y=0, w=0.4, h=0.2)
    boxes, _ = platform_node_rects(nodes, edges, rect)
    assert max(box.bottom for box in boxes.values()) <= rect.bottom + 1e-6


def test_sample_has_no_overflow() -> None:
    issues = validate_deck(load_financial_tech_sample())
    assert not issues, [(issue.code, issue.message) for issue in issues]


def test_financial_export_is_editable_and_complete() -> None:
    deck = load_financial_tech_sample()
    data = render_deck_to_pptx(deck)
    report = verify_pptx(data, deck)
    assert report.passed, report.issues
    data.seek(0)
    ppt = Presentation(data)
    names = [shape.name for slide in ppt.slides for shape in slide.shapes]
    assert sum(name.startswith("financial-chevron-") for name in names) == 4
    assert sum(name.startswith("financial-platform-node-") for name in names) == 12
    headers = [shape for shape in ppt.slides[0].shapes if shape.name.startswith("financial-chevron-")]
    for header in headers:
        assert header.fill.type == MSO_FILL_TYPE.GRADIENT
        assert header.fill.gradient_angle == 270
        assert len(header.fill.gradient_stops) == 2
        assert header.fill.gradient_stops[0].color.rgb != header.fill.gradient_stops[1].color.rgb
    assert not any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in ppt.slides for shape in slide.shapes
    )
    title = next(
        shape
        for shape in ppt.slides[0].shapes
        if shape.has_text_frame and shape.text == "企业家客户服务矩阵"
    )
    assert title.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


@pytest.mark.parametrize("theme_id", ["financial-tech", "enterprise-dark", "enterprise"])
def test_existing_sample_can_switch_theme(theme_id: str) -> None:
    deck = load_sample_deck()
    report = verify_pptx(render_deck_to_pptx(deck, theme_id), deck)
    assert report.passed, report.issues


@pytest.mark.asyncio
async def test_preview_download_exports_edited_text() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        payload = deepcopy(load_financial_tech_sample().model_dump())
        payload["slides"][0]["blocks"][0]["text"] = "自定义金融服务"
        response = await client.post(
            "/api/v1/design/themes/financial-tech/preview.pptx", json=payload
        )
        assert response.status_code == 200
        from io import BytesIO

        ppt = Presentation(BytesIO(response.content))
        assert any(
            shape.has_text_frame and shape.text == "自定义金融服务"
            for shape in ppt.slides[0].shapes
        )
        assert (await client.get("/api/v1/design/themes/nope/preview.pptx")).status_code == 404
