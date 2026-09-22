"""Enterprise planning, source fidelity, composition and export regression checks."""

import uuid
from io import BytesIO

import pytest
from pptx import Presentation

from app.domain.content import (
    CardsBlock,
    ChartBlock,
    Deck,
    DiagramBlock,
    Slide,
    TableBlock,
    TextBlock,
)
from app.domain.enterprise_layout import bind_planned_evidence, compose_report_slide
from app.domain.enterprise_quality import check_planned_data, check_report_quality
from app.domain.evidence import comparable_series, evidence_problem, prepare_page_plan
from app.domain.flex_layout import FlexContainer, FlexLeaf, iter_leaf_block_ids
from app.domain.mermaid import ensure_mermaid, mermaid_for_diagram
from app.domain.outline import DeckBlueprint, EvidenceItem, OutlinePage, ReportBrief
from app.domain.theme import get_theme
from app.render.pptx import render_deck_to_pptx
from app.render.verify import verify_pptx
from app.services.preview import load_preview_status, preview_fingerprint, save_preview_status


def fact(value="120", period="Q1", **changes):
    return EvidenceItem.model_validate(
        {
            "source_ref": "S1:1",
            "quote": f"{period}集团收入{value}万元。",
            "metric": "收入",
            "value": value,
            "unit": "万元",
            "period": period,
            "scope": "集团",
            **changes,
        }
    )


def category_fact(category: str, value: str):
    return fact(
        value=value,
        quote=f"Q1集团收入中{category}{value}万元。",
        period="Q1",
        category=category,
    )


def plan(**changes):
    return OutlinePage.model_validate(
        {
            "title": "收入提升",
            "objective": "说明经营变化",
            "key_points": ["结果", "行动"],
            "layout_id": "chart",
            "evidence_kind": "trend",
            "narrative_role": "performance",
            "key_message": "收入提升",
            "evidence": [fact(), fact("140", "Q2")],
            **changes,
        }
    )


def slide():
    return Slide(
        id="s1",
        layout_id="bullets",
        layout_mode="flex",
        blocks=[
            TextBlock(id="kicker", slot_id="kicker", text="经营复盘"),
            TextBlock(id="heading-xyz", slot_id="heading-xyz", text="收入提升"),
            TextBlock(id="insight", slot_id="insight", text="持续观察经营质量"),
        ],
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[
                FlexLeaf(id="l1", block_id="kicker", text_style="caption"),
                FlexLeaf(id="l2", block_id="heading-xyz", text_style="title"),
                FlexLeaf(id="l3", block_id="insight", text_style="body"),
            ],
        ),
    )


@pytest.mark.parametrize(
    "change",
    [
        {"value": "12"},
        {"value": "120%"},
        {"unit": "%"},
        {"metric": "利润"},
        {"period": "Q2"},
        {"scope": "子公司"},
        {"quote": "材料不存在的句子"},
        {"value": "NaN"},
        {"value": "Infinity"},
    ],
)
def test_false_evidence_rejected(change):
    assert evidence_problem(fact(**change), {"S1:1": "Q1集团收入120万元。"})


def test_valid_evidence_and_incompatible_series():
    assert evidence_problem(fact(), {"S1:1": "Q1集团收入120万元。"}) is None
    assert len(comparable_series([fact(), fact("140", "Q2")])) == 2
    assert not comparable_series([fact(), fact("140", "Q2", scope="子公司")])
    assert not comparable_series([fact(), fact("140")])


def test_missing_data_downgrades_without_inventing_metrics():
    result = prepare_page_plan(plan(), {})
    assert result.evidence_kind == "narrative"
    assert result.layout_id == "bullets"
    assert result.evidence == []
    assert result.planning_notes


def test_flow_without_numeric_evidence_keeps_diagram_intent():
    result = prepare_page_plan(
        plan(
            evidence_kind="flow",
            visual_type="flow",
            key_points=["输入：识别问题", "处理：拆解任务", "输出：形成方案"],
            evidence=[],
        ),
        {},
    )

    assert result.evidence_kind == "flow"
    assert result.visual_type == "flow"


def test_diagram_blocks_have_mermaid_source_for_rendering():
    page = plan(
        evidence_kind="flow",
        visual_type="flow",
        key_points=["输入：识别问题", "处理：拆解任务", "输出：形成方案"],
        evidence=[],
    )
    result = bind_planned_evidence(slide(), page)
    diagram = next(block for block in result.blocks if block.type == "diagram")

    assert diagram.mermaid is not None
    assert diagram.mermaid.startswith("flowchart TB")
    assert "-->" in diagram.mermaid


def test_incident_flow_generates_branch_and_merge_mermaid():
    page = plan(
        evidence_kind="flow",
        visual_type="flow",
        title="故障上报与处理流程",
        key_points=["故障预警", "上报主管部门负责人"],
        evidence=[],
    )
    source = slide().model_copy(
        update={
            "blocks": [
                *slide().blocks,
                CardsBlock(
                    id="incident-cards",
                    slot_id="cards",
                    items=[
                        {"title": "故障预警", "desc": ""},
                        {"title": "上报主管部门负责人", "desc": ""},
                        {"title": "APP故障", "desc": ""},
                        {"title": "ERP系统故障", "desc": ""},
                        {"title": "主机/网络故障", "desc": ""},
                        {"title": "数据库故障", "desc": ""},
                        {"title": "成立联合保障团队", "desc": "紧急排查处理"},
                        {"title": "系统恢复", "desc": ""},
                        {"title": "复盘总结", "desc": ""},
                    ],
                ),
            ]
        }
    )
    result = bind_planned_evidence(source, page)
    diagram = next(block for block in result.blocks if block.type == "diagram")
    pairs = {(edge.source, edge.target) for edge in diagram.edges}

    assert diagram.mermaid is not None
    assert diagram.mermaid.startswith("flowchart TB")
    assert ("step-2", "step-3") in pairs
    assert ("step-2", "step-6") in pairs
    assert ("step-3", "step-7") in pairs
    assert ("step-6", "step-7") in pairs
    assert ("step-7", "step-8") in pairs


def test_authored_mermaid_is_preserved():
    nodes = [
        {"id": "a", "title": "开始"},
        {"id": "b", "title": "结束"},
    ]
    edges = [{"source": "a", "target": "b", "label": "下一步"}]
    authored = "flowchart LR\n  A[开始] --> B[结束]"

    block = DiagramBlock(
        id="diagram",
        slot_id="visual",
        diagram_type="flow",
        mermaid=authored,
        nodes=nodes,
        edges=edges,
    )

    assert ensure_mermaid(block.mermaid, block.diagram_type, block.nodes, block.edges) == authored
    assert mermaid_for_diagram(block.diagram_type, block.nodes, block.edges).startswith(
        "flowchart TB"
    )


def test_composition_keeps_each_block_once_and_title_style():
    result = compose_report_slide(bind_planned_evidence(slide(), plan()), plan())
    assert result.layout_tree.children[0].block_id == "heading-xyz"
    ids = iter_leaf_block_ids(result.layout_tree)
    assert len(ids) == len(set(ids)) == len(result.blocks)
    assert set(ids) == {b.id for b in result.blocks}
    assert not check_planned_data(result, plan())
    chart = next(b for b in result.blocks if b.type == "chart")
    chart.series[0].values[1] = 999
    assert check_planned_data(result, plan())[0].severity == "error"


def test_composition_evidence_defaults_to_pie_chart():
    page = plan(
        evidence_kind="composition",
        visual_type="auto",
        evidence=[category_fact("线上", "60"), category_fact("线下", "40")],
    )
    result = bind_planned_evidence(slide(), page)
    chart = next(b for b in result.blocks if b.type == "chart")

    assert chart.chart_type == "pie"
    assert chart.categories == ["线上", "线下"]
    assert chart.series[0].values == [60.0, 40.0]


def test_comparison_evidence_defaults_to_bar_chart():
    page = plan(
        evidence_kind="comparison",
        visual_type="auto",
        evidence=[category_fact("A方案", "120"), category_fact("B方案", "95")],
    )
    result = bind_planned_evidence(slide(), page)
    chart = next(b for b in result.blocks if b.type == "chart")

    assert chart.chart_type == "bar"
    assert chart.categories == ["A方案", "B方案"]


def test_flow_plan_binds_native_diagram_block():
    page = plan(
        evidence_kind="flow",
        visual_type="flow",
        key_points=["采集：汇总业务输入", "校验：核对口径", "输出：形成汇报"],
        evidence=[],
    )
    result = bind_planned_evidence(slide(), page)
    diagram = next(b for b in result.blocks if b.type == "diagram")

    assert isinstance(diagram, DiagramBlock)
    assert diagram.diagram_type == "flow"
    assert [node.title for node in diagram.nodes] == ["采集", "校验", "输出"]
    assert [(edge.source, edge.target) for edge in diagram.edges] == [
        ("step-1", "step-2"),
        ("step-2", "step-3"),
    ]


def test_flow_visual_type_overrides_narrative_evidence_kind():
    page = plan(
        evidence_kind="narrative",
        visual_type="flow",
        key_points=["识别问题", "拆解任务"],
        evidence=[],
    )

    result = bind_planned_evidence(slide(), page)
    diagram = next(block for block in result.blocks if block.type == "diagram")

    assert isinstance(diagram, DiagramBlock)
    assert diagram.diagram_type == "flow"


def test_timeline_reuses_table_rows_as_diagram_nodes():
    page = plan(
        evidence_kind="timeline",
        visual_type="timeline",
        key_points=["方向确认", "试点验证"],
        evidence=[],
    )
    source = slide().model_copy(
        update={
            "blocks": [
                TextBlock(id="title", slot_id="title", text="阶段路线"),
                TableBlock(
                    id="table",
                    slot_id="table",
                    header=["阶段", "时间", "交付物"],
                    rows=[
                        ["方向确认", "2026Q3", "明确目标与评价指标"],
                        ["试点验证", "2026Q4", "完成首批试点交付"],
                    ],
                ),
            ]
        }
    )

    result = bind_planned_evidence(source, page)
    diagram = next(block for block in result.blocks if block.type == "diagram")

    assert isinstance(diagram, DiagramBlock)
    assert [node.title for node in diagram.nodes] == ["方向确认", "试点验证"]
    assert "时间：2026Q3" in diagram.nodes[0].desc
    assert "交付物：明确目标与评价指标" in diagram.nodes[0].desc


def test_qualitative_driver_page_gets_semantic_diagram_fallback():
    page = plan(
        title="背后驱动力：开发方式正在变化",
        evidence_kind="narrative",
        narrative_role="driver",
        key_points=["工具普及：降低重复劳动", "需求变化：业务理解更重要", "协作升级：人机共同交付"],
        evidence=[],
    )
    result = bind_planned_evidence(slide(), page)
    diagram = next(b for b in result.blocks if b.type == "diagram")

    assert diagram.diagram_type == "flow"
    assert [node.title for node in diagram.nodes] == ["工具普及", "需求变化", "协作升级"]


def test_flow_page_is_full_width_and_drops_competing_structures():
    page = plan(
        evidence_kind="flow",
        visual_type="flow",
        narrative_role="supporting",
        key_points=["输入：识别问题", "处理：拆解任务", "输出：形成方案"],
        evidence=[],
    )
    blocks = [
        TextBlock(id="title", slot_id="title", text="流程"),
        DiagramBlock(
            id="diagram",
            slot_id="visual",
            diagram_type="flow",
            nodes=[
                {"id": "n1", "title": "输入", "desc": "识别问题"},
                {"id": "n2", "title": "处理", "desc": "拆解任务"},
                {"id": "n3", "title": "输出", "desc": "形成方案"},
            ],
        ),
        CardsBlock(
            id="cards",
            slot_id="cards",
            items=[{"title": "阶段信息", "desc": "不应与流程图同页堆叠"}],
        ),
        TableBlock(id="table", slot_id="table", header=["阶段"], rows=[["输入"]]),
    ]
    source = Slide(
        id="flow-layout",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )

    result = compose_report_slide(source, page)

    assert [block.type for block in result.blocks] == ["text", "diagram"]
    stage = result.layout_tree.children[1]
    assert isinstance(stage, FlexContainer)
    assert stage.type == "overlay"
    visual = stage.children[0]
    assert isinstance(visual, FlexLeaf)
    assert visual.block_id == "diagram"
    assert visual.bleed
    assert all(
        not (hasattr(child, "id") and child.id == "evidence-and-insight")
        for child in result.layout_tree.children
    )


def test_flow_page_keeps_one_compact_support_text_without_squeezing():
    page = plan(
        evidence_kind="flow",
        visual_type="flow",
        narrative_role="supporting",
        key_points=["输入：识别问题", "处理：拆解任务", "输出：形成方案"],
        evidence=[],
    )
    blocks = [
        TextBlock(id="title", slot_id="title", text="流程"),
        DiagramBlock(
            id="diagram",
            slot_id="visual",
            diagram_type="flow",
            nodes=[
                {"id": "n1", "title": "输入", "desc": "识别问题"},
                {"id": "n2", "title": "处理", "desc": "拆解任务"},
                {"id": "n3", "title": "输出", "desc": "形成方案"},
            ],
        ),
        TextBlock(
            id="support",
            slot_id="support",
            text="通过标准化流程把问题拆解为输入、处理和输出三个阶段，帮助团队稳定交付。",
        ),
        TextBlock(
            id="extra",
            slot_id="extra",
            text="这段较长的补充说明不应与流程图一起堆叠。",
        ),
        CardsBlock(
            id="cards",
            slot_id="cards",
            items=[{"title": "阶段信息", "desc": "节点已经承载阶段信息"}],
        ),
    ]
    source = Slide(
        id="flow-support",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[
                FlexLeaf(id=f"leaf-{block.id}", block_id=block.id)
                for block in blocks
            ],
        ),
    )

    result = compose_report_slide(source, page)

    assert [block.id for block in result.blocks] == ["title", "diagram", "support"]
    stage = result.layout_tree.children[1]
    assert isinstance(stage, FlexContainer)
    assert stage.type == "overlay"
    assert iter_leaf_block_ids(stage) == ["diagram", "support"]
    overlay = stage.children[1]
    assert isinstance(overlay, FlexContainer)
    assert overlay.children[0].id.startswith("spacer-")


def test_chart_page_drops_table_and_cards():
    page = plan(evidence_kind="comparison", visual_type="bar")
    blocks = [
        TextBlock(id="title", slot_id="title", text="方案比较"),
        ChartBlock(
            id="chart",
            slot_id="visual",
            chart_type="bar",
            categories=["A", "B"],
            series=[{"name": "准备度", "values": [80, 60]}],
        ),
        TextBlock(id="insight", slot_id="insight", text="A 方案领先，差距需要解释"),
        CardsBlock(
            id="cards",
            slot_id="cards",
            items=[{"title": "不保留", "desc": "图表页不再叠加卡片"}],
        ),
        TableBlock(id="table", slot_id="table", header=["方案"], rows=[["A"]]),
    ]
    source = Slide(
        id="chart-layout",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )

    result = compose_report_slide(source, page)

    assert [block.type for block in result.blocks] == ["text", "chart", "text"]
    stage = result.layout_tree.children[1]
    assert isinstance(stage, FlexContainer)
    assert stage.type == "overlay"
    assert iter_leaf_block_ids(stage) == ["chart", "insight"]


def test_chart_page_moves_long_support_to_speaker_notes_instead_of_fourth_band():
    page = plan(evidence_kind="comparison", visual_type="bar")
    long_support = (
        "决策请求：确认下一阶段推进方向与资源安排。建议按照明确目标、盘点现状、"
        "设计方案、试点验证、规模推广五步推进，并对应 2026Q3 方向确认、"
        "2026Q4 方案设计、2027Q1 试点运行、2027Q2 加速推广的节奏。"
    )
    blocks = [
        TextBlock(id="title", slot_id="title", text="方案比较"),
        ChartBlock(
            id="chart",
            slot_id="visual",
            chart_type="bar",
            categories=["A", "B"],
            series=[{"name": "准备度", "values": [80, 60]}],
        ),
        TextBlock(id="support", slot_id="support", text=long_support),
    ]
    source = Slide(
        id="chart-long-support",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )

    result = compose_report_slide(source, page)

    assert [block.id for block in result.blocks] == ["title", "chart"]
    assert result.layout_tree.children[0].block_id == "title"
    stage = result.layout_tree.children[1]
    assert isinstance(stage, FlexContainer)
    assert stage.type == "overlay"
    assert iter_leaf_block_ids(stage) == ["chart"]
    assert result.speaker_notes is not None
    assert "未放入画布的补充说明" in result.speaker_notes
    assert "决策请求" in result.speaker_notes


def test_action_page_chooses_cards_or_table_once():
    page = plan(
        evidence_kind="actions",
        narrative_role="action",
        evidence=[],
        key_points=["明确动作", "指定负责人"],
    )
    blocks = [
        TextBlock(id="title", slot_id="title", text="行动计划"),
        CardsBlock(
            id="cards",
            slot_id="cards",
            items=[
                {"title": "推进", "desc": "本周完成"},
                {"title": "复盘", "desc": "下周完成"},
            ],
        ),
        TableBlock(
            id="table",
            slot_id="table",
            header=["行动", "负责人"],
            rows=[["推进", "项目组"]],
        ),
        TextBlock(id="summary", slot_id="summary", text="优先推进第一项"),
    ]
    source = Slide(
        id="action-layout",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )

    result = compose_report_slide(source, page)
    structured = [block for block in result.blocks if block.type in {"cards", "table"}]

    assert len(structured) == 1
    assert all(
        block.type not in {"cards", "table"}
        for block in result.blocks
        if block is not structured[0]
    )


def test_summary_cards_and_bullets_stack_instead_of_compressing_into_a_row():
    page = plan(
        evidence_kind="narrative",
        narrative_role="executive_summary",
        evidence=[],
        key_points=["判断", "下一步"],
    )
    blocks = [
        TextBlock(id="title", slot_id="title", text="执行摘要"),
        CardsBlock(
            id="cards",
            slot_id="cards",
            items=[
                {"title": "方向", "desc": "围绕核心问题推进"},
                {"title": "资源", "desc": "优先保障关键环节"},
                {"title": "节奏", "desc": "按阶段验收交付"},
            ],
        ),
        TextBlock(id="summary", slot_id="summary", text="先完成方向确认，再进入试点"),
    ]
    source = Slide(
        id="summary-layout",
        layout_id="bullets",
        layout_mode="flex",
        blocks=blocks,
        layout_tree=FlexContainer(
            type="column",
            id="root",
            children=[FlexLeaf(id=f"leaf-{block.id}", block_id=block.id) for block in blocks],
        ),
    )

    result = compose_report_slide(source, page)
    comparison = next(
        (
            child
            for child in result.layout_tree.children
            if getattr(child, "id", None) == "comparison"
        ),
        None,
    )

    assert comparison is None


def test_unplanned_chart_rejected_in_fixed_mode_too():
    data = ChartBlock(
        id="chart",
        slot_id="visual",
        chart_type="line",
        categories=["Q1"],
        series=[{"name": "收入", "values": [999]}],
        unit="万元",
    )
    result = Slide(id="fixed", layout_id="chart", blocks=[data])
    assert check_planned_data(result, plan(evidence=[]))


def test_pptx_roundtrip_checks_actual_numeric_values():
    result = compose_report_slide(bind_planned_evidence(slide(), plan()), plan())
    deck = Deck(id="d", title="报告", theme_id="enterprise", slides=[result])
    payload = render_deck_to_pptx(deck).getvalue()
    assert verify_pptx(payload, deck).passed
    chart = next(b for b in result.blocks if b.type == "chart")
    chart.series[0].values[0] = 999
    assert not verify_pptx(payload, deck).passed
    assert any(shape.has_chart for shape in Presentation(BytesIO(payload)).slides[0].shapes)


def test_deck_missing_summary_and_decision_are_visible():
    slides = [slide().model_copy(update={"id": str(i)}) for i in range(5)]
    plans = {s.id: plan(evidence=[]) for s in slides}
    deck = Deck(id="d", title="报告", theme_id="enterprise", slides=slides)
    codes = {
        i.code
        for i in check_report_quality(
            deck,
            plans,
            {},
            DeckBlueprint(),
            ReportBrief(scenario="business_review", decision_request="确认推广方向"),
        )
    }
    assert {"missing_summary", "missing_decision", "missing_close"} <= codes


def test_preview_fingerprint_changes_with_content_and_theme():
    deck = Deck(id="d", title="报告", theme_id="enterprise", slides=[slide()])
    original = preview_fingerprint(deck, get_theme("enterprise"))
    assert original == preview_fingerprint(deck, get_theme("enterprise"))
    assert original != preview_fingerprint(deck, get_theme("enterprise-dark"))
    deck.slides[0].blocks[0].text = "最新经营复盘"
    assert original != preview_fingerprint(deck, get_theme("enterprise"))


def test_preview_timeout_is_retryable(monkeypatch):
    key = f"previews/test/{uuid.uuid4()}"
    save_preview_status(key, "queued")
    assert load_preview_status(key)["status"] == "queued"
    monkeypatch.setattr("app.services.preview.time.time", lambda: 10**12)
    assert load_preview_status(key)["status"] == "failed"
