from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from app.domain.content import Deck, DiagramBlock, FinancialTableBlock, Slide, TableBlock
from app.domain.diagram_layout import layout_diagram, parse_mermaid
from app.domain.diagram_style import themed_flowchart_style
from app.domain.enterprise_layout import bind_planned_evidence, compose_report_slide
from app.domain.flex_layout import FlexContainer, FlexLeaf
from app.domain.outline import OutlinePageDraft
from app.domain.table_layout import layout_table
from app.domain.table_pagination import split_table_slide
from app.domain.text_metrics import wrap_text
from app.domain.theme import get_theme
from app.domain.visual_layout import bind_source_diagram
from app.ingest.markdown import MarkdownParser
from app.llm.base import SlideGenerationInput
from app.render.pptx import render_deck_to_pptx
from app.render.verify import verify_pptx
from app.workflows.slide import prepare_slide_input

SOURCE = Path(__file__).parent / "fixtures/music_execution.mmd"


@pytest.mark.parametrize("theme", ["enterprise", "financial-tech"])
def test_native_flowchart_uses_reference_style_in_every_theme(theme):
    block = DiagramBlock(
        id="g",
        slot_id="g",
        diagram_type="flow",
        nodes=[
            {"id": "a", "title": "检查条件", "kind": "decision"},
            {"id": "b", "title": "继续处理", "desc": "保留完整说明"},
            {"id": "c", "title": "退回修改"},
        ],
        edges=[
            {"source": "a", "target": "b", "label": "通过"},
            {"source": "a", "target": "c", "label": "未通过"},
        ],
    )
    slide = Slide(
        id="s",
        layout_id="bullets",
        layout_mode="flex",
        blocks=[block],
        layout_tree=FlexContainer(
            id="root", type="column", children=[FlexLeaf(id="g", block_id="g")]
        ),
    )
    # A stale saved layout must be rebuilt using the shared style and final slot.
    block.geometry = layout_diagram(block, 600, 300).model_copy(update={"version": 1})
    data = render_deck_to_pptx(Deck(id="d", title="test", theme_id=theme, slides=[slide]))
    shapes = Presentation(BytesIO(data.getvalue())).slides[0].shapes
    colors = themed_flowchart_style(get_theme(theme))
    nodes = {shape.name.rsplit("-", 1)[-1]: shape for shape in shapes if "-node-" in shape.name}
    assert str(nodes["b"].fill.fore_color.rgb) == colors["node_fill"].lstrip("#")
    assert str(nodes["a"].fill.fore_color.rgb) == colors["background"].lstrip("#")
    assert nodes["a"]._element.xpath(".//a:prstDash[@val='dash']")
    routes = [shape for shape in shapes if shape.name.startswith("diagram-edge-")]
    assert len(routes) == 2
    assert all(str(shape.line.color.rgb) == colors["connector"].lstrip("#") for shape in routes)
    assert any(shape._element.xpath(".//a:quadBezTo") for shape in routes)
    assert all(shape._element.xpath(".//a:tailEnd[@type='arrow']") for shape in routes)
    text_colors = [
        str(run.font.color.rgb)
        for shape in shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    assert text_colors and set(text_colors) == {colors["text"].lstrip("#")}


def assert_attached_routes(block, geometry):
    nodes = {node.id: node for node in geometry.nodes}

    def attached(point, node):
        x, y = point
        assert node.x - 0.01 <= x <= node.x + node.width + 0.01
        assert node.y - 0.01 <= y <= node.y + node.height + 0.01
        assert (
            min(
                abs(x - node.x),
                abs(x - node.x - node.width),
                abs(y - node.y),
                abs(y - node.y - node.height),
            )
            < 0.01
        )

    for edge, route in zip(block.edges, geometry.edges, strict=True):
        attached(route.points[0], nodes[edge.source])
        attached(route.points[-1], nodes[edge.target])


def test_source_flow_survives_excerpt_binding_composition_and_native_export():
    source = SOURCE.read_text(encoding="utf-8")
    section = (
        MarkdownParser().parse(f"# 执行流程\n```mermaid\n{source}```\n说明".encode()).sections[0]
    )
    payload = SlideGenerationInput(
        deck_title="Music",
        tone="专业",
        position=1,
        total_pages=1,
        page_title="执行流程",
        objective="完整流程",
        key_points=["个人库优先", "允许时联网"],
        layout_id="bullets",
        sections=[{"ref": "S1:1", **section.model_dump()}],
    )
    payload = prepare_slide_input(payload, max_section_chars=10)
    assert payload.sections[0].diagrams[0].source.strip() == source.strip()
    page = OutlinePageDraft(
        title="执行流程",
        objective="完整流程",
        key_points=["个人库优先", "允许时联网"],
        layout_id="bullets",
        evidence_kind="flow",
        visual_type="flow",
    )
    slide = Slide(id="flow", layout_id="bullets", layout_mode="flex", blocks=[])
    slide = bind_source_diagram(slide, payload.sections, requested=True)
    slide = bind_planned_evidence(slide, page)
    slide = compose_report_slide(slide, page)
    block = next(b for b in slide.blocks if isinstance(b, DiagramBlock))
    assert len(block.nodes) == 12
    assert len(block.edges) == 13
    assert {(n.id, n.kind) for n in block.nodes if n.kind == "decision"} == {
        ("F", "decision"),
        ("G", "decision"),
    }
    assert ("H", "D", None) in {(e.source, e.target, e.label) for e in block.edges}
    assert {e.label for e in block.edges if e.label} == {"满足", "不足", "允许", "不允许或预算耗尽"}
    g = layout_diagram(block, 856, 360)
    assert g.font_size >= 12
    assert len(g.edges) == len(block.edges)
    assert_attached_routes(block, g)
    for i, node in enumerate(g.nodes):
        assert 0 <= node.x <= g.width - node.width + 0.01
        assert 0 <= node.y <= g.height - node.height + 0.01
        for other in g.nodes[i + 1 :]:
            assert (
                node.x + node.width <= other.x
                or other.x + other.width <= node.x
                or node.y + node.height <= other.y
                or other.y + other.height <= node.y
            )
    for edge in g.edges:
        for a, b in zip(edge.points, edge.points[1:], strict=False):
            assert a[0] == pytest.approx(b[0]) or a[1] == pytest.approx(b[1])
    block.geometry = g
    slide = Slide(
        id="flow",
        layout_id="bullets",
        layout_mode="flex",
        blocks=[block],
        layout_tree=FlexContainer(
            id="root", type="column", children=[FlexLeaf(id="graph", block_id=block.id)]
        ),
    )
    # Use the exact slot dimensions rather than scaling the cached geometry.
    block.geometry = layout_diagram(block, 856, 454)
    deck = Deck(id="test", title="Flow", theme_id="financial-tech", slides=[slide])
    data = render_deck_to_pptx(deck)
    assert verify_pptx(data, deck).passed
    data.seek(0)
    ppt = Presentation(data)
    assert not any(shape.shape_type == 13 for shape in ppt.slides[0].shapes)
    assert any(shape._element.xpath(".//a:tailEnd") for shape in ppt.slides[0].shapes)


@pytest.mark.parametrize("direction", ["TB", "LR", "BT", "RL"])
def test_wrapped_branches_and_cycles_stay_attached_in_each_direction(direction):
    nodes, edges, _ = parse_mermaid(SOURCE.read_text(encoding="utf-8"))
    block = DiagramBlock(
        id="graph",
        slot_id="graph",
        diagram_type="flow",
        nodes=nodes,
        edges=edges,
        direction=direction,
    )
    assert_attached_routes(block, layout_diagram(block, 856, 360))


def test_parser_preserves_bidirectional_edges_and_legacy_boundaries():
    from app.ingest.models import SourceSection

    section = SourceSection(
        level=1,
        heading="Architecture",
        locator="S1:1",
        text='flowchart LR\n A["Agent"] <--> B["Model"]\n图示后的说明文字',
    )
    nodes, edges, direction = parse_mermaid(section.diagrams[0].source)
    assert direction == "LR"
    assert {n.title for n in nodes} == {"Agent", "Model"}
    assert {(e.source, e.target) for e in edges} == {("A", "B"), ("B", "A")}


def test_table_keeps_long_identifiers_and_repeats_headers_without_losing_rows():
    identifier = "music_recommendation_runs_" * 4
    block = FinancialTableBlock(
        id="table",
        slot_id="table",
        columns=[identifier, "music_track_features"],
        unit="表结构对照",
        rows=[{"label": f"字段 {i}", "values": [identifier, "完整信息" * 8]} for i in range(12)],
    )
    g = layout_table(block, 856, 454)
    assert g.column_widths[0] < g.column_widths[1]
    assert g.font_size >= 12
    assert "".join(g.cells[0][1].split()) == identifier
    block.geometry = g
    slide = Slide(
        id="table",
        layout_id="bullets",
        layout_mode="flex",
        blocks=[block],
        layout_tree=FlexContainer(
            id="root", type="column", children=[FlexLeaf(id="table-leaf", block_id="table")]
        ),
    )
    chunks = split_table_slide(slide)
    assert len(chunks) > 1
    assert [r for chunk in chunks for r in chunk.blocks[0].rows] == block.rows
    assert all(chunk.blocks[0].geometry.cells[0] == g.cells[0] for chunk in chunks)
    assert all(chunk.blocks[0].geometry.height <= 454 for chunk in chunks)


def test_one_oversized_row_is_an_explicit_error():
    block = TableBlock(id="table", slot_id="table", header=["字段"], rows=[["内容" * 2000]])
    block.geometry = layout_table(block, 856, 454)
    slide = Slide(
        id="table",
        layout_id="bullets",
        layout_mode="flex",
        blocks=[block],
        layout_tree=FlexContainer(
            id="root", type="column", children=[FlexLeaf(id="leaf", block_id="table")]
        ),
    )
    with pytest.raises(ValueError, match="row exceeds"):
        split_table_slide(slide)


def test_markdown_tables_preserve_cell_boundaries_and_source_positions():
    parsed = MarkdownParser().parse(
        b"# Schema\n| Field | Type |\n|---|---|\n| run_id | CHAR(32) |\n"
    )
    table = parsed.sections[0].tables[0]
    assert table.header == ["Field", "Type"]
    assert table.rows == [["run_id", "CHAR(32)"]]
    assert table.locator
    wrapped = wrap_text("a" * 90, width_pt=70, size_pt=14)
    assert "".join(wrapped) == "a" * 90
    assert len(wrapped) > 1
