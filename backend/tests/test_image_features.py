import asyncio
import json
import uuid
from unittest.mock import AsyncMock

import httpx
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from redis.exceptions import RedisError
from sqlalchemy import select

from app.api.deps import get_queue
from app.core.config import get_settings
from app.core.db import async_session_factory
from app.domain.content import Deck, ImageBlock, Slide, TextBlock
from app.domain.enterprise_layout import bind_planned_evidence, compose_report_slide
from app.domain.evidence import prepare_page_plan
from app.domain.flex_layout import iter_leaf_block_ids
from app.domain.image_planning import bind_page_image, plan_page_image, search_queries
from app.domain.outline import ImagePlan, OutlinePageDraft
from app.images.bailian import BailianImageProvider
from app.images.base import ImageAsset, ImageCandidate, ImageRequest
from app.images.errors import ImageFetchError
from app.images.pipeline import ImagePipeline
from app.images.unsplash import UnsplashImageProvider
from app.main import app
from app.models.project import ProjectOutline
from app.models.slide import Slide as SlideRow
from app.render.pptx import render_deck_to_pptx
from app.services.image_candidates import load_candidate
from app.services.media import media_url, store_image
from app.worker.deck_tasks import generate_deck
from app.worker.image_tasks import _save_result, generate_slide_images
from tests.test_decks_api import FakeSlideGenerator, _confirmed_project
from tests.test_images import MINIMAL_PNG, _project_with_image_slide, _sign_up


@pytest.fixture
async def client():
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                                 base_url="http://test") as value:
        yield value


def page(**kwargs):
    return OutlinePageDraft(
        title="故宫游览",
        objective="介绍景点与游览建议",
        key_points=["沿中轴线参观", "预约以官方信息为准"],
        layout_id="bullets",
        **kwargs,
    )


def candidate():
    return ImageCandidate(
        id="forbidden-city",
        url="https://images.unsplash.com/photo-demo",
        thumbnail_url="https://images.unsplash.com/photo-demo?w=300",
        description="Forbidden City, Beijing",
        width=1800,
        height=1200,
        credit="Photo by Ada on Unsplash",
        credit_url="https://unsplash.com/p/demo",
    )


def image(blocks):
    return next(block for block in blocks if block.get("type") == "image")


class Queue:
    def __init__(self, *, fail=False):
        self.calls = []
        self.fail = fail

    async def enqueue_job(self, *args, **kwargs):
        if self.fail:
            raise RedisError("offline")
        self.calls.append(args)
        return object()


class Provider:
    source = "stock"

    def __init__(self, *, fail=False):
        self.calls = 0
        self.fail = fail

    def available(self):
        return True

    async def fetch(self, request):
        self.calls += 1
        return (
            None
            if self.fail
            else ImageAsset(
                data=MINIMAL_PNG, content_type="image/png", source="stock", asset_id="new-photo"
            )
        )


def test_travel_planning_preserves_photo_through_evidence_and_layout():
    planned = plan_page_image(page(), deck_title="春节北京旅游")
    assert planned.visual_type == "photo"
    assert planned.image_plan.require_real
    assert planned.image_plan.queries[0] == "Beijing Forbidden City"
    prepared = prepare_page_plan(planned, {}, topic_mode=True)
    assert prepared.visual_type == "photo"
    slide = bind_page_image(
        Slide(
            id="s",
            layout_id="bullets",
            layout_mode="flex",
            blocks=[TextBlock(id="title", slot_id="title", text="故宫")],
        ),
        prepared,
    )
    slide = bind_planned_evidence(slide, prepared, topic_mode=True)
    slide = compose_report_slide(slide, prepared, layout_template="visual_left")
    images = [b for b in slide.blocks if b.type == "image"]
    assert len(images) == 1
    assert images[0].id in iter_leaf_block_ids(slide.layout_tree)


def test_explicit_illustration_overrides_travel_photo_default():
    planned = plan_page_image(page(visual_type="illustration"), deck_title="北京旅游")
    assert planned.image_plan.source == "generated"
    assert not planned.image_plan.require_real


def test_photo_plan_does_not_replace_explicit_timeline():
    planned = plan_page_image(
        page(visual_type="timeline", evidence_kind="timeline"), deck_title="北京旅游"
    )
    assert planned.image_plan is None
    assert planned.visual_type == "timeline"


def test_generic_photo_prefers_stock_and_allows_ai_fallback():
    planned = plan_page_image(OutlinePageDraft(
        title="AI 算力基础设施", objective="介绍计算机行业", key_points=["算力需求增长", "建设机房"],
        layout_id="image-right",
        visual_type="photo", visual="数据中心机房内 GPU 服务器阵列与蓝色灯光",
    ))
    assert planned.image_plan.source == "stock"
    assert not planned.image_plan.require_real
    assert planned.image_plan.queries[0] == "server rack"


@pytest.mark.parametrize("subject", ["故宫建筑", "iPhone 17 产品实拍"])
def test_specific_photo_remains_real_even_when_plan_relaxes_requirement(subject):
    planned = plan_page_image(OutlinePageDraft(
        title="对象介绍", objective="介绍图片主体", key_points=["展示外观", "介绍特点"],
        layout_id="image-right",
        visual_type="photo", image_plan=ImagePlan(subject=subject, require_real=False),
    ))
    assert planned.image_plan.source == "stock" and planned.image_plan.require_real


def test_search_queries_use_concrete_objects_and_preserve_product_identity():
    assert search_queries("AI PC 笔记本电脑与 AI 手机并排展示的产品实拍")[0] == "laptop smartphone"
    assert search_queries("iPhone 17 手机", ["iPhone 17 product photo"])[0] == "iPhone 17 product photo"


@pytest.mark.parametrize(("subject", "description", "matched"), [
    ("数据中心机房内 GPU 服务器阵列与蓝色灯光", "Rows of servers in a data centre", True),
    ("数据中心机房内 GPU 服务器阵列与蓝色灯光", "A blue coat on a clothing rack", False),
    ("AI PC 笔记本电脑与 AI 手机并排展示的产品实拍", "A notebook computer and cell phone", True),
    ("AI PC 笔记本电脑与 AI 手机并排展示的产品实拍", "A laptop on an office desk", False),
    ("iPhone 17 手机", "iPhone 16 smartphone product photo", False),
    ("iPhone 17 手机", "iPhone 17 smartphone product photo", True),
    ("服务器机房", "3D rendered server racks", False),
])
async def test_stock_match_checks_core_objects_and_specific_identity(subject, description, matched):
    def handler(request):
        return httpx.Response(200, json={"results": [{
            "id": "photo", "description": description, "width": 1800, "height": 1200,
            "urls": {"regular": "https://images.unsplash.com/photo"},
        }]})

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        candidates = await UnsplashImageProvider(client=transport, access_key="test").search(
            ImageRequest(prompt="", query=subject, subject=subject, aspect_ratio=1.5,
                         preferred_source="stock", queries=["smartphone"] if "iPhone" in subject else [])
        )
    assert candidates[0].metadata_matched is matched


async def test_stock_search_time_budget_falls_back_to_ai():
    async def handler(request):
        await asyncio.sleep(0.1)
        return httpx.Response(200, json={"results": []})

    generated = Provider()
    generated.source = "generated"
    generated.fetch = AsyncMock(return_value=ImageAsset(
        data=MINIMAL_PNG, content_type="image/png", source="generated",
    ))
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        stock = UnsplashImageProvider(client=transport, access_key="test", search_timeout_seconds=0.01)
        result = await ImagePipeline([generated, stock]).fetch(ImageRequest(
            prompt="p", query="server rack", aspect_ratio=1.5, preferred_source="stock",
        ))
    assert result.source == "generated" and generated.fetch.await_count == 1


async def test_stock_search_keeps_partial_candidates_on_timeout():
    calls = 0

    async def handler(request):
        nonlocal calls
        calls += 1
        if calls > 1:
            await asyncio.sleep(0.1)
        return httpx.Response(200, json={"results": [{
            "id": "photo", "description": "office desk",
            "urls": {"regular": "https://images.unsplash.com/photo"},
        }]})

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        candidates = await UnsplashImageProvider(
            client=transport, access_key="test", search_timeout_seconds=0.01,
        ).search(ImageRequest(prompt="", query="server rack", aspect_ratio=1.5, preferred_source="stock"))
    assert len(candidates) == 1 and not candidates[0].metadata_matched


async def test_wanx_submits_once_and_polls_correct_endpoint(monkeypatch):
    submissions = []
    polls = []
    monkeypatch.setattr("app.images.bailian.asyncio.sleep", AsyncMock())

    def handler(request):
        if request.method == "POST":
            submissions.append(json.loads(request.content))
            assert request.headers["X-DashScope-Async"] == "enable"
            assert request.url.path.endswith("/text2image/image-synthesis")
            return httpx.Response(200, json={"output": {"task_id": "task-1"}})
        if request.url.path.endswith("/tasks/task-1"):
            polls.append(request)
            if len(polls) == 1:
                return httpx.Response(503)
            return httpx.Response(
                200,
                json={
                    "output": {
                        "task_status": "SUCCEEDED",
                        "results": [{"url": "https://cdn.test/a.png"}],
                    }
                },
            )
        return httpx.Response(200, content=MINIMAL_PNG)

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        asset = await BailianImageProvider(client=transport, api_key="test", model="wanx-v1").fetch(
            ImageRequest(prompt="AI illustration", query="AI", aspect_ratio=1.777)
        )
    assert asset.data == MINIMAL_PNG
    assert len(submissions) == 1 and len(polls) == 2
    assert submissions[0]["input"] == {"prompt": "AI illustration"}
    assert submissions[0]["parameters"]["size"] == "1280*720"


async def test_wanx_timeout_does_not_resubmit():
    submissions = []

    def handler(request):
        if request.method == "POST":
            submissions.append(request)
            return httpx.Response(200, json={"output": {"task_id": "task-1"}})
        return httpx.Response(200, json={"output": {"task_status": "RUNNING"}})

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        asset = await BailianImageProvider(
            client=transport, api_key="test", model="wanx-v1", timeout_seconds=0.02
        ).fetch(ImageRequest(prompt="p", query="q", aspect_ratio=1))
    assert asset is None and len(submissions) == 1


async def test_real_photo_rejects_unrelated_first_result_and_deduplicates():
    downloads = []

    def handler(request):
        if request.url.path.endswith("/search/photos"):
            assert request.url.params["per_page"] == "8"
            return httpx.Response(
                200,
                json={
                    "results": [
                        {
                            "id": "wrong",
                            "description": "A building in Paris",
                            "urls": {"regular": "https://images.unsplash.com/wrong"},
                            "user": {},
                        },
                        {
                            "id": "correct",
                            "description": "Forbidden City Beijing",
                            "urls": {"regular": "https://images.unsplash.com/correct"},
                            "user": {},
                        },
                    ]
                },
            )
        downloads.append(request.url.path)
        return httpx.Response(200, content=MINIMAL_PNG)

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        provider = UnsplashImageProvider(client=transport, access_key="test")
        request = ImageRequest(
            prompt="", query="故宫", subject="故宫", aspect_ratio=1.5, require_real=True
        )
        asset = await provider.fetch(request)
        missing = await provider.fetch(request.model_copy(update={"excluded_ids": {"correct"}}))
    assert asset.asset_id == "correct" and missing is None
    assert downloads == ["/correct"]


async def test_real_photo_does_not_fall_back_to_ai():
    stock, generated = Provider(fail=True), Provider()
    generated.source = "generated"
    result = await ImagePipeline([generated, stock]).fetch(
        ImageRequest(prompt="", query="故宫", aspect_ratio=1, require_real=True)
    )
    assert result is None and generated.calls == 0


async def test_landmark_search_prefers_building_description_over_portrait():
    def handler(_request):
        return httpx.Response(200, json={"results": [
            {"id": "guard", "description": "A palace guard at Forbidden City Beijing",
             "width": 1800, "height": 1200,
             "urls": {"regular": "https://images.unsplash.com/guard"}, "user": {}},
            {"id": "building", "description": "Hall of Supreme Harmony in Forbidden City Beijing",
             "width": 2400, "height": 1700,
             "urls": {"regular": "https://images.unsplash.com/building"}, "user": {}},
        ]})

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        candidates = await UnsplashImageProvider(client=transport, access_key="test").search(
            ImageRequest(prompt="", query="故宫", subject="故宫", aspect_ratio=1.5)
        )
    assert candidates[0].id == "building"
    assert all(item.metadata_matched for item in candidates)
    assert all("尚未" in item.match_reason for item in candidates)


async def test_real_landmark_search_relaxes_orientation_when_no_metadata_match():
    orientations = []

    def handler(request):
        orientations.append(request.url.params.get("orientation"))
        if "orientation" in request.url.params:
            return httpx.Response(200, json={"results": [
                {"id": "unrelated", "description": "An office building",
                 "urls": {"regular": "https://images.unsplash.com/office"}, "user": {}},
            ]})
        return httpx.Response(200, json={"results": [
            {"id": "landmark", "description": "Forbidden City Beijing palace",
             "width": 2400, "height": 1600,
             "urls": {"regular": "https://images.unsplash.com/palace"}, "user": {}},
        ]})

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
        candidates = await UnsplashImageProvider(client=transport, access_key="test").search(
            ImageRequest(prompt="", query="故宫", subject="故宫", aspect_ratio=0.6,
                         require_real=True)
        )
    assert candidates[0].id == "landmark" and candidates[0].metadata_matched
    assert orientations[-1] is None and len(orientations) <= 5


@pytest.fixture
def queue():
    value = Queue()
    app.dependency_overrides[get_queue] = lambda: value
    yield value
    app.dependency_overrides.pop(get_queue, None)


async def test_search_cache_apply_and_foreign_candidate_rejection(client, monkeypatch):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    search = AsyncMock(return_value=[candidate()])
    monkeypatch.setattr(UnsplashImageProvider, "search", search)
    monkeypatch.setattr(
        UnsplashImageProvider,
        "download",
        AsyncMock(
            return_value=ImageAsset(
                data=MINIMAL_PNG,
                content_type="image/png",
                source="stock",
                credit=candidate().credit,
                credit_url=candidate().credit_url,
                asset_id=candidate().id,
            )
        ),
    )
    path = f"/api/v1/projects/{project['id']}/deck"
    result = await client.post(f"{path}/images/search", headers=headers, json={"query": "故宫"})
    assert result.status_code == 200
    await client.post(f"{path}/images/search", headers=headers, json={"query": "故宫"})
    assert search.await_count == 1
    body = {
        "revision": slide.revision,
        "search_id": result.json()["search_id"],
        "candidate_id": candidate().id,
    }
    ok = await client.post(
        f"{path}/slides/{slide.id}/blocks/img-1/image/apply", headers=headers, json=body
    )
    assert ok.status_code == 200
    selected = image(ok.json()["blocks"])
    assert selected["locked"] and selected["image_status"] == "ready"
    assert selected["credit_url"] == candidate().credit_url
    conflict = await client.post(
        f"{path}/slides/{slide.id}/blocks/img-1/image/apply", headers=headers, json=body
    )
    assert conflict.status_code == 409
    with pytest.raises(Exception, match="过期"):
        await load_candidate(uuid.uuid4(), body["search_id"], candidate().id)


async def test_image_job_completes_without_regenerating_text(client, queue):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    path = f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image/generate"
    response = await client.post(
        path, headers=headers, json={"revision": 1, "source": "stock", "query": "故宫"}
    )
    assert response.status_code == 200 and response.json()["status"] == "ready"
    block = image(response.json()["blocks"])
    assert block["image_status"] == "queued"
    provider = Provider()
    await generate_slide_images({"image_pipeline": ImagePipeline([provider])}, *queue.calls[0][1:])
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert row.status == "ready" and image(row.blocks)["image_status"] == "ready"
        assert image(row.blocks)["image_asset_id"] == "new-photo"
        assert row.revision == response.json()["revision"] + 1


async def test_upload_invalidates_running_image_job(client, queue):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    path = f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image"
    generated = await client.post(f"{path}/generate", headers=headers, json={"revision": 1})
    upload = await client.put(
        path,
        headers=headers,
        data={"revision": str(generated.json()["revision"])},
        files={"file": ("photo.png", MINIMAL_PNG, "image/png")},
    )
    assert upload.status_code == 200
    provider = Provider()
    await generate_slide_images({"image_pipeline": ImagePipeline([provider])}, *queue.calls[0][1:])
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert image(row.blocks)["source"] == "upload" and image(row.blocks)["locked"]
        assert image(row.blocks)["url"] == image(upload.json()["blocks"])["url"]
    assert provider.calls == 0


async def test_stale_result_clears_job_without_overwriting_user_edit(client, queue):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    path = f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image/generate"
    generated = await client.post(path, headers=headers, json={"revision": 1})
    block = image(generated.json()["blocks"])
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        row.revision += 1
        row.blocks = [{**b, "alt": "User edit"} if b.get("type") == "image" else b
                      for b in row.blocks]
        await session.commit()
    await _save_result(
        uuid.UUID(project["id"]),
        slide.id,
        block["image_job_id"],
        generated.json()["revision"],
        {"img-1": {**block, "alt": "Stale"}},
        None,
    )
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert image(row.blocks)["alt"] == "User edit"
        assert image(row.blocks)["image_status"] == "failed"
        assert image(row.blocks)["image_job_id"] is None


async def test_image_queue_failure_is_visible_and_retryable(client, queue):
    queue.fail = True
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    path = f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image/generate"
    result = await client.post(path, headers=headers, json={"revision": 1})
    assert result.status_code == 503
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert row.status == "ready" and image(row.blocks)["image_status"] == "failed"
        assert image(row.blocks)["image_job_id"] is None


async def test_add_image_converts_fixed_layout_and_preserves_content(client):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        row.layout_id = "bullets"
        row.blocks = [
            {"id": "title", "slot_id": "title", "type": "text", "text": "Keep title"},
            {"id": "body", "slot_id": "body", "type": "bullets", "items": ["Keep body"]},
        ]
        await session.commit()
    result = await client.post(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/images",
        headers=headers,
        json={"revision": 1, "subject": "故宫"},
    )
    assert result.status_code == 200
    body = result.json()
    assert body["layout_mode"] == "flex" and len(body["blocks"]) == 3
    assert body["blocks"][0]["text"] == "Keep title"
    assert body["blocks"][1]["items"] == ["Keep body"]


async def test_lock_blocks_generation_until_unlocked(client, queue):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers, locked=True)
    path = f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image"
    rejected = await client.post(f"{path}/generate", headers=headers, json={"revision": 1})
    assert rejected.status_code == 409
    unlocked = await client.patch(
        f"{path}/lock", headers=headers, json={"revision": 1, "locked": False}
    )
    assert unlocked.status_code == 200
    accepted = await client.post(
        f"{path}/generate", headers=headers, json={"revision": unlocked.json()["revision"]}
    )
    assert accepted.status_code == 200


async def test_failed_replacement_keeps_old_image_and_public_error(client, queue):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    old_url = "/media/previous.png"
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        row.blocks = [{**b, "url": old_url, "source": "stock"}
                      if b.get("type") == "image" else b for b in row.blocks]
        await session.commit()
    response = await client.post(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image/generate",
        headers=headers, json={"revision": 1, "source": "stock"},
    )
    assert response.status_code == 200
    provider = Provider()
    provider.fetch = AsyncMock(side_effect=ImageFetchError("图库请求额度已用完，请稍后重试"))
    await generate_slide_images({"image_pipeline": ImagePipeline([provider])}, *queue.calls[0][1:])
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        block = image(row.blocks)
        assert row.status == "ready" and row.error is None
        assert block["url"] == old_url and block["image_status"] == "failed"
        assert "额度" in block["image_error"] and block["image_job_id"] is None


async def test_concurrent_image_workers_reserve_different_stock_assets(client, queue, monkeypatch):
    headers = await _sign_up(client)
    project, first = await _project_with_image_slide(client, headers)
    async with async_session_factory() as session:
        second = SlideRow(
            project_id=first.project_id, outline_page_id=uuid.uuid4(), position=2,
            layout_id=first.layout_id, layout_mode="fixed", title="Second page",
            status="ready", blocks=first.blocks, issues=[], revision=1,
        )
        session.add(second)
        await session.commit()
        await session.refresh(second)
    monkeypatch.setattr(get_settings(), "unsplash_access_key", "test")
    matched = candidate().model_copy(update={"metadata_matched": True})
    candidates = [matched, matched.model_copy(update={"id": "second-photo"})]
    monkeypatch.setattr(UnsplashImageProvider, "search", AsyncMock(return_value=candidates))

    async def download(_self, selected):
        await asyncio.sleep(0)
        return ImageAsset(data=MINIMAL_PNG, content_type="image/png",
                          source="stock", asset_id=selected.id)

    monkeypatch.setattr(UnsplashImageProvider, "download", download)
    for row in [first, second]:
        response = await client.post(
            f"/api/v1/projects/{project['id']}/deck/slides/{row.id}/blocks/img-1/image/generate",
            headers=headers, json={"revision": 1, "source": "stock", "query": "故宫"},
        )
        assert response.status_code == 200
    await asyncio.gather(*(generate_slide_images({}, *call[1:]) for call in queue.calls))
    async with async_session_factory() as session:
        rows = [await session.get(SlideRow, row.id) for row in [first, second]]
        assert all(image(row.blocks)["image_status"] == "ready" for row in rows)
        assert {image(row.blocks)["image_asset_id"] for row in rows} == {
            "forbidden-city", "second-photo"
        }


async def test_apply_rechecks_revision_after_remote_download(client, monkeypatch):
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    path = f"/api/v1/projects/{project['id']}/deck"
    monkeypatch.setattr(UnsplashImageProvider, "search", AsyncMock(return_value=[candidate()]))

    async def download(*_args):
        async with async_session_factory() as session:
            row = await session.get(SlideRow, slide.id)
            row.revision += 1
            row.title = "User changed title during download"
            await session.commit()
        return ImageAsset(data=MINIMAL_PNG, content_type="image/png", source="stock")

    monkeypatch.setattr(UnsplashImageProvider, "download", download)
    found = await client.post(f"{path}/images/search", headers=headers, json={"query": "故宫"})
    result = await client.post(
        f"{path}/slides/{slide.id}/blocks/img-1/image/apply", headers=headers,
        json={"revision": 1, "search_id": found.json()["search_id"],
              "candidate_id": candidate().id},
    )
    assert result.status_code == 409
    async with async_session_factory() as session:
        row = await session.get(SlideRow, slide.id)
        assert row.title == "User changed title during download"
        assert image(row.blocks)["url"] is None


async def test_deck_queues_images_only_after_body_ready(client, queue):
    headers = await _sign_up(client)
    project = await _confirmed_project(client, headers, layout_mode="flex")
    async with async_session_factory() as session:
        outline = await session.scalar(select(ProjectOutline).where(
            ProjectOutline.project_id == uuid.UUID(project["id"])
        ))
        outline.pages = [{**p, "visual_type": "photo", "image_plan": {
            "subject": "故宫", "queries": ["Beijing Forbidden City"],
            "source": "stock", "require_real": True,
        }} if index == 1 else p for index, p in enumerate(outline.pages)]
        await session.commit()
    started = await client.post(
        f"/api/v1/projects/{project['id']}/deck/generate", headers=headers, json={},
    )
    assert started.status_code == 202
    ids = queue.calls[0][2]
    observed = []
    original_enqueue = queue.enqueue_job

    async def enqueue(*args, **kwargs):
        if args[0] == "generate_slide_images":
            async with async_session_factory() as session:
                row = await session.get(SlideRow, uuid.UUID(args[2]))
                observed.append((row.status, row.blocks))
        return await original_enqueue(*args, **kwargs)

    queue.enqueue_job = enqueue
    await generate_deck(
        {"slide_generator": FakeSlideGenerator(), "redis": queue}, project["id"], ids
    )
    assert observed and all(status == "ready" for status, _ in observed)
    planned = [b for _, blocks in observed for b in blocks
               if b.get("type") == "image" and (b.get("image_plan") or {}).get("require_real")]
    assert planned and planned[0]["image_plan"]["queries"] == ["Beijing Forbidden City"]
    assert planned[0]["image_status"] == "queued"


def test_export_keeps_stock_picture_and_attribution():
    key = store_image(user_id=uuid.uuid4(), project_id=uuid.uuid4(),
                      data=MINIMAL_PNG, extension=".png")
    deck = Deck(id="deck", title="故宫", theme_id="midnight", slides=[Slide(
        id="slide", layout_id="image-right", speaker_notes="讲解故宫",
        blocks=[TextBlock(id="title", slot_id="title", text="故宫"), ImageBlock(
            id="photo", slot_id="image", alt="故宫", source="stock", url=media_url(key),
            credit=candidate().credit, credit_url=candidate().credit_url,
        )],
    )])
    presentation = Presentation(render_deck_to_pptx(deck))
    page = presentation.slides[0]
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in page.shapes)
    notes = page.notes_slide.notes_text_frame.text
    assert "讲解故宫" in notes and candidate().credit in notes and candidate().credit_url in notes
